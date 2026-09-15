import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Color Palette - Modern Tech Dark Mode Accent
    DARK_BG = RGBColor(15, 23, 42)      # #0f172a Deep Slate
    CARD_BG = RGBColor(30, 41, 59)      # #1e293b Card Blue
    ACCENT_PURPLE = RGBColor(168, 85, 247) # #a855f7 Vibrant Purple
    ACCENT_BLUE = RGBColor(59, 130, 246)  # #3b82f6 Bright Blue
    TEXT_WHITE = RGBColor(248, 250, 252) # #f8fafc White
    TEXT_MUTED = RGBColor(148, 163, 184) # #94a3b8 Muted Grey
    ACCENT_GREEN = RGBColor(34, 197, 94)  # #22c55e Emerald

    def add_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = DARK_BG
        bg.line.fill.background()

    def add_header(slide, title_text, category_text="CSE PATHFINDER"):
        # Header category
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(10), Inches(0.4))
        tf_cat = cat_box.text_frame
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = ACCENT_PURPLE

        # Header main title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.5), Inches(0.8))
        tf_title = title_box.text_frame
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(28)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_WHITE

    def add_card(slide, left, top, width, height, title, points, accent_color=ACCENT_BLUE):
        # Card Background
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = accent_color
        card.line.width = Pt(1.5)

        # Text Frame
        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_right = Inches(0.3)
        tf.margin_top = Inches(0.3)
        tf.margin_bottom = Inches(0.3)

        # Title
        p0 = tf.paragraphs[0]
        p0.text = title
        p0.font.size = Pt(20)
        p0.font.bold = True
        p0.font.color.rgb = TEXT_WHITE
        p0.space_after = Pt(12)

        # Points
        for pt_text in points:
            p = tf.add_paragraph()
            p.text = "• " + pt_text
            p.font.size = Pt(14)
            p.font.color.rgb = TEXT_MUTED
            p.space_after = Pt(8)

    # ----------------------------------------------------
    # SLIDE 1: TITLE SLIDE
    # ----------------------------------------------------
    blank_layout = prs.slide_layouts[6]
    slide1 = prs.slides.add_slide(blank_layout)
    add_background(slide1)

    # Decorative Accent Bar
    bar = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(2.2), Inches(0.15), Inches(3.2))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT_PURPLE
    bar.line.fill.background()

    # Main Title Box
    t_box = slide1.shapes.add_textbox(Inches(1.4), Inches(2.0), Inches(10.5), Inches(3.5))
    tf1 = t_box.text_frame
    tf1.word_wrap = True

    p1 = tf1.paragraphs[0]
    p1.text = "CSE PathFinder"
    p1.font.size = Pt(54)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_WHITE
    p1.space_after = Pt(10)

    p2 = tf1.add_paragraph()
    p2.text = "Intelligent Career Trajectory & 8-Semester Guidance Platform"
    p2.font.size = Pt(24)
    p2.font.bold = True
    p2.font.color.rgb = ACCENT_BLUE
    p2.space_after = Pt(20)

    p3 = tf1.add_paragraph()
    p3.text = "Multi-Track Diagnostic Quiz  •  Dynamic Inclination Radar  •  Java 21 Backend  •  Supabase Cloud DB"
    p3.font.size = Pt(14)
    p3.font.color.rgb = TEXT_MUTED

    # ----------------------------------------------------
    # SLIDE 2: PROBLEM STATEMENT & MOTIVATION
    # ----------------------------------------------------
    slide2 = prs.slides.add_slide(blank_layout)
    add_background(slide2)
    add_header(slide2, "Problem Statement & Industry Context", "THE CHALLENGE")

    add_card(slide2, Inches(0.8), Inches(1.8), Inches(3.6), Inches(4.8), 
             "Domain Paralysis", 
             ["CSE students are overwhelmed by vast domain options (AI, Web, Cyber, Cloud, IoT, Robotics, VLSI, Networking).",
              "Lack of early domain clarity leads to fragmented learning and weak project portfolios."], 
             ACCENT_PURPLE)

    add_card(slide2, Inches(4.8), Inches(1.8), Inches(3.6), Inches(4.8), 
             "Generic Roadmaps", 
             ["Traditional advice provides static, one-size-fits-all roadmaps without considering student strengths.",
              "Students lack semester-by-semester structured skill targets and practical mini-projects."], 
             ACCENT_BLUE)

    add_card(slide2, Inches(8.8), Inches(1.8), Inches(3.6), Inches(4.8), 
             "Progress Tracking Gap", 
             ["No unified system to track progress across 8 semesters alongside diagnostic assessment history.",
              "Need for a robust 24/7 web platform with cloud persistence and local fallback."], 
             ACCENT_GREEN)

    # ----------------------------------------------------
    # SLIDE 3: PROPOSED SOLUTION
    # ----------------------------------------------------
    slide3 = prs.slides.add_slide(blank_layout)
    add_background(slide3)
    add_header(slide3, "The CSE PathFinder Solution", "THE SOLUTION")

    add_card(slide3, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8), 
             "🎯 Multi-Track Diagnostic Engine", 
             ["Custom 5-stage quiz runner evaluating General, Software, and Hardware inclinations.",
              "Generates normalized affinity scores across 9 specialized Computer Science domains.",
              "Calculates cross-domain affinity correlation matrix for accurate trajectory mapping."], 
             ACCENT_PURPLE)

    add_card(slide3, Inches(6.8), Inches(1.8), Inches(5.6), Inches(4.8), 
             "🗺️ 8-Semester Roadmap & Tracker", 
             ["Tailored 8-semester curriculum for every CSE domain with mini-projects & capstone prep.",
              "Interactive subtopic checklist with instant progress persistence.",
              "Personalized notes & external resource link logger for each course."], 
             ACCENT_BLUE)

    # ----------------------------------------------------
    # SLIDE 4: SYSTEM ARCHITECTURE & TECH STACK
    # ----------------------------------------------------
    slide4 = prs.slides.add_slide(blank_layout)
    add_background(slide4)
    add_header(slide4, "Full-Stack System Architecture", "TECHNICAL STACK")

    add_card(slide4, Inches(0.8), Inches(1.8), Inches(3.6), Inches(4.8), 
             "Frontend Layer", 
             ["React 19 & Vite",
              "Tailwind CSS Dark Theme",
              "Lucide-React Icons",
              "Chart.js & React-ChartJS-2",
              "React Router SPA Routing"], 
             ACCENT_PURPLE)

    add_card(slide4, Inches(4.8), Inches(1.8), Inches(3.6), Inches(4.8), 
             "Java 21 Backend Server", 
             ["OpenJDK 21 LTS Engine",
              "com.sun.net.httpserver API",
              "Native SHA-256 Hashing",
              "Multi-Threaded Cached Pool",
              "CORS Allowed Origin Headers"], 
             ACCENT_BLUE)

    add_card(slide4, Inches(8.8), Inches(1.8), Inches(3.6), Inches(4.8), 
             "Database & Cloud Layer", 
             ["SQLite JDBC (database.db)",
              "Supabase Cloud PostgreSQL",
              "apiService Dual DB Layer",
              "LocalStorage Local Fallback",
              "Vercel Serverless Ready"], 
             ACCENT_GREEN)

    # ----------------------------------------------------
    # SLIDE 5: KEY FEATURES & WORKFLOW
    # ----------------------------------------------------
    slide5 = prs.slides.add_slide(blank_layout)
    add_background(slide5)
    add_header(slide5, "User Workflow & Core Platform Features", "PLATFORM FEATURES")

    add_card(slide5, Inches(0.8), Inches(1.8), Inches(5.6), Inches(2.2), 
             "1. Student Registration & Auth", 
             ["Secure registration with SHA-256 hashed passwords.",
              "Captures name, college, year, and branch profile info."], 
             ACCENT_PURPLE)

    add_card(slide5, Inches(6.8), Inches(1.8), Inches(5.6), Inches(2.2), 
             "2. Interactive Quiz Runner", 
             ["5-stage stepper with progress indicators.",
              "Supports 7 multi-choice options A through G."], 
             ACCENT_BLUE)

    add_card(slide5, Inches(0.8), Inches(4.3), Inches(5.6), Inches(2.3), 
             "3. Inclination Analytics & Radar", 
             ["Visualizes top recommended domain match percentage.",
              "Renders Chart.js radar chart comparing domain scores."], 
             ACCENT_GREEN)

    add_card(slide5, Inches(6.8), Inches(4.3), Inches(5.6), Inches(2.3), 
             "4. Semester Roadmap & Persistence", 
             ["8-semester curriculum breakdown with subtopic lists.",
              "Saves status (Not Started/In Progress/Completed) & notes."], 
             ACCENT_PURPLE)

    # ----------------------------------------------------
    # SLIDE 6: JAVA BACKEND API ENDPOINTS
    # ----------------------------------------------------
    slide6 = prs.slides.add_slide(blank_layout)
    add_background(slide6)
    add_header(slide6, "Java 21 REST API Endpoint Specifications", "BACKEND ARCHITECTURE")

    add_card(slide6, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8), 
             "Authentication & User APIs", 
             ["GET /api/health — Health check & server status",
              "POST /api/register — Inserts new student record",
              "POST /api/login — Validates credentials & returns student JSON profile"], 
             ACCENT_BLUE)

    add_card(slide6, Inches(6.8), Inches(1.8), Inches(5.6), Inches(4.8), 
             "Quiz & Progress APIs", 
             ["POST /api/submit-result — Stores quiz result & JSON scores",
              "GET /api/results/{email} — Retrieves user result history",
              "POST /api/progress — Updates semester status & notes",
              "GET /api/progress/{email} — Returns progress records"], 
             ACCENT_GREEN)

    # ----------------------------------------------------
    # SLIDE 7: RESULTS & TESTING VERIFICATION
    # ----------------------------------------------------
    slide7 = prs.slides.add_slide(blank_layout)
    add_background(slide7)
    add_header(slide7, "Testing, Verification & Deployment", "VERIFICATION")

    add_card(slide7, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8), 
             "Localhost Verification", 
             ["Automated browser subagent verified registration & login on http://localhost:5175.",
              "Java REST API server verified with 100% success rate (200 OK across all 7 endpoints).",
              "Vite production build verified in 2.50s with 0 compilation errors."], 
             ACCENT_PURPLE)

    add_card(slide7, Inches(6.8), Inches(1.8), Inches(5.6), Inches(4.8), 
             "Zero Downtime Architecture", 
             ["Dual Database Service (apiService.js) ensures app works 100% locally even without internet.",
              "Vercel + Supabase Cloud PostgreSQL readiness for 24/7 online deployment.",
              "Single 1-click startup script (start_project.bat) launches Java backend & React frontend."], 
             ACCENT_GREEN)

    # ----------------------------------------------------
    # SLIDE 8: CONCLUSION & FUTURE SCOPE
    # ----------------------------------------------------
    slide8 = prs.slides.add_slide(blank_layout)
    add_background(slide8)
    add_header(slide8, "Future Scope & Conclusion", "CONCLUSION")

    add_card(slide8, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8), 
             "Future Scope", 
             ["AI Mentor Integration (LLM-based personalized career advice chatbot).",
              "Real-time Job Market Skill Alignment (scraping tech stack demands).",
              "Peer Mentorship & GitHub Portfolio Review Hub."], 
             ACCENT_BLUE)

    add_card(slide8, Inches(6.8), Inches(1.8), Inches(5.6), Inches(4.8), 
             "Conclusion", 
             ["CSE PathFinder bridges the gap between academic learning and industry readiness.",
              "Empowers CSE students with personalized domain diagnosis and actionable 8-semester roadmaps.",
              "Robust full-stack Java 21 + React 19 architecture ready for local & cloud deployment."], 
             ACCENT_PURPLE)

    output_path = "CSE_PathFinder_Presentation.pptx"
    prs.save(output_path)
    print(f"[SUCCESS] Presentation generated successfully: {output_path}")

if __name__ == "__main__":
    create_presentation()
